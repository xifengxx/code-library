#!/usr/bin/env python3
"""build-word-fancy.py — python-docx 精装版 Word：封面 + 目录 + 分类章节表格。

数据源：web/src/data/repos.json（由 build-index.mjs 生成），与网页端共享同一份数据。
依赖：pip3 install python-docx
"""
import json
import sys
from pathlib import Path

try:
    from docx import Document
    from docx.enum.section import WD_SECTION
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Inches, Pt, RGBColor
except ImportError:
    print("❌ 未安装 python-docx：pip3 install python-docx")
    sys.exit(1)

ROOT = Path(__file__).resolve().parent.parent
DATA_PATH = ROOT / "web" / "src" / "data" / "repos.json"
OUT_PATH = ROOT / "output" / "word" / "代码仓库参照库-精装版.docx"

ACCENT = RGBColor(0x22, 0x55, 0x99)
GRAY = RGBColor(0x66, 0x66, 0x66)

CATEGORIES = ['ai-llm', 'frontend', 'backend', 'database', 'devops', 'tools',
              'mobile', 'learning', 'templates', 'misc']
LABELS = {
    'ai-llm': 'AI / LLM', 'frontend': '前端', 'backend': '后端', 'database': '数据库 / 存储',
    'devops': 'DevOps / 云', 'tools': '开发工具 / CLI', 'mobile': '移动端',
    'learning': '学习资源', 'templates': '项目模板 / 示例', 'misc': '其他',
}

if not DATA_PATH.exists():
    print("❌ 缺少 web/src/data/repos.json，请先运行 node scripts/build-index.mjs")
    sys.exit(1)

data = json.loads(DATA_PATH.read_text(encoding="utf-8"))
entries = data["entries"]
by_cat = {c: [e for e in entries if e["category"] == c] for c in CATEGORIES}
total = len(entries)

doc = Document()

# 默认字体
style = doc.styles["Normal"]
style.font.name = "Microsoft YaHei"
style.font.size = Pt(10.5)


def para(text="", size=10.5, bold=False, color=None, align=None, space_after=6):
    p = doc.add_paragraph()
    if align is not None:
        p.alignment = align
    p.paragraph_format.space_after = Pt(space_after)
    r = p.add_run(text)
    r.font.size = Pt(size)
    r.bold = bold
    if color:
        r.font.color.rgb = color
    return p


# ── 封面 ────────────────────────────────────────────────
para("", size=18, space_after=0)
para("📚 代码仓库参照库", size=28, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=12)
para("个人代码仓库收藏 / 索引 / 实践手册", size=14, color=GRAY, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=24)
para(f"共收录 {total} 个仓库 · 生成日期 {__import__('datetime').date.today().isoformat()}",
     size=11, color=GRAY, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=36)
para("分类一览：", size=12, bold=True)
for c in CATEGORIES:
    n = len(by_cat[c])
    para(f"  {LABELS[c]}（{n}）", size=11)

doc.add_section(WD_SECTION.NEW_PAGE)

# ── 目录（手列分类，正文为章节标题） ─────────────────────
para("目录", size=20, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=12)
for c in CATEGORIES:
    if by_cat[c]:
        para(f"{LABELS[c]} · {len(by_cat[c])} 个", size=12, bold=True, space_after=4)

# ── 每个分类一个章节 + 表格 ──────────────────────────────
for c in CATEGORIES:
    list_e = by_cat[c]
    if not list_e:
        continue
    doc.add_section(WD_SECTION.NEW_PAGE)
    para(f"{LABELS[c]}（{len(list_e)}）", size=18, bold=True, color=ACCENT, space_after=10)

    table = doc.add_table(rows=1, cols=5)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    hdr = table.rows[0].cells
    for i, h in enumerate(["仓库", "一句话", "语言", "Stars", "标签"]):
        hdr[i].text = ""
        run = hdr[i].paragraphs[0].add_run(h)
        run.bold = True
        run.font.color.rgb = ACCENT

    for e in list_e:
        row = table.add_row().cells
        name = row[0].paragraphs[0].add_run(e["name"])
        name.bold = True
        name.font.color.rgb = ACCENT
        if e.get("repo_url"):
            row[0].paragraphs[0].add_run("\n" + e["repo_url"]).font.size = Pt(8)
        row[1].text = (e.get("sections", {}).get("一句话", "") or "")[:80]
        row[2].text = e.get("language", "") or ""
        stars = e.get("stars")
        row[3].text = f"{stars/1000:.1f}k" if stars and stars >= 1000 else (str(stars) if stars else "")
        row[4].text = " ".join("#" + t for t in (e.get("tags") or []))

OUT_PATH.parent.mkdir(parents=True, exist_ok=True)
doc.save(str(OUT_PATH))
print(f"🎉 精装版已生成 → {OUT_PATH}")
